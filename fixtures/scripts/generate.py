"""Generates feature fixtures with fake data.

Usage:
    .venv/bin/python fixtures/scripts/generate.py           # feature set
    .venv/bin/python fixtures/scripts/generate.py --large   # also the 100k-row perf file
"""

import datetime as dt
import io
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import (
    CellIsRule,
    ColorScaleRule,
    DataBarRule,
    FormulaRule,
    IconSetRule,
    Rule,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle

OUT = Path(__file__).resolve().parent.parent / "generated"
random.seed(42)


def save(wb: Workbook, name: str) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    wb.properties.creator = "xl-viewer fixtures"
    wb.properties.lastModifiedBy = "xl-viewer fixtures"
    path = OUT / name
    wb.save(path)
    print(path.relative_to(OUT.parent.parent))


def basic() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Basic"
    ws.append(["Item", "Qty", "Price", "In stock", "Updated", "Note"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")
    items = ["Apple", "Banana", "Cherry", "Date", "Elderberry", "Fig", "Grape"]
    for i, name in enumerate(items):
        ws.append([name, random.randint(1, 500), round(random.uniform(0.5, 30), 2), i % 2 == 0, dt.date(2026, 1, 1) + dt.timedelta(days=i * 9), None])
    ws["F3"] = "Long text that overflows into the empty cells to the right"
    ws["A10"] = "Theme colours"
    for i in range(6):
        c = ws.cell(row=11, column=i + 1, value=f"accent{i + 1}")
        c.fill = PatternFill("solid", fgColor=["4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"][i])
    ws["A13"] = "Bold"
    ws["A13"].font = Font(bold=True)
    ws["B13"] = "Italic"
    ws["B13"].font = Font(italic=True)
    ws["C13"] = "Underline"
    ws["C13"].font = Font(underline="single")
    ws["D13"] = "Strike"
    ws["D13"].font = Font(strike=True)
    ws["E13"] = "Red 14pt"
    ws["E13"].font = Font(color="FF0000", size=14)
    ws["F13"] = "#DIV/0!"
    ws["A14"] = "=1/0"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["G"].hidden = True
    ws["G2"] = "hidden column"
    ws.row_dimensions[8].hidden = True
    ws.row_dimensions[13].height = 24
    for col in "BC":
        for row in range(2, 9):
            ws[f"{col}{row}"].number_format = "#,##0" if col == "B" else "#,##0.00"
    for row in range(2, 9):
        ws[f"E{row}"].number_format = "yyyy-mm-dd"

    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = "This sheet is hidden"
    hidden.sheet_state = "hidden"
    second = wb.create_sheet("한글 시트")
    second["A1"] = "한글 텍스트와 English 섞기"
    second["A2"] = 1234567.891
    second.column_dimensions["A"].width = 30
    save(wb, "basic.xlsx")


def number_formats() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formats"
    ws.append(["Format", "Value", "Shown"])
    rows = [
        ("General", 1234567.891),
        ("General", 0.000012345),
        ("General", 123456789012345),
        ("0", 3.7),
        ("0.00", 3.14159),
        ("#,##0", -1234567),
        ("#,##0.00;[Red]-#,##0.00", -42.5),
        ("0%", 0.256),
        ("0.00%", 0.12345),
        ("0.00E+00", 123456),
        ("# ?/?", 1.75),
        ('"₩"#,##0', 15000),
        ('_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-', 0),
        ("yyyy-mm-dd", dt.date(2026, 10, 28)),
        ("mmmm d, yyyy", dt.date(2026, 10, 28)),
        ('yyyy"년" m"월" d"일"', dt.date(2026, 10, 28)),
        ("ddd", dt.date(2026, 10, 28)),
        ("h:mm AM/PM", dt.time(15, 30)),
        ("[h]:mm:ss", 1.5),
        ("@", "text"),
        ('[Blue][>100]"high";[Red][<0]"neg";"mid"', 150),
        ('[Blue][>100]"high";[Red][<0]"neg";"mid"', -1),
    ]
    for fmt, value in rows:
        ws.append([fmt, value, value])
        r = ws.max_row
        ws.cell(row=r, column=3).number_format = fmt
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 22
    ws["E2"] = 123456789
    ws["E2"].number_format = "#,##0"
    ws["E3"] = dt.date(2026, 10, 28)
    ws["E3"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["E"].width = 5
    ws["F1"] = "Narrow column E shows ###"
    save(wb, "number-formats.xlsx")


def borders() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Borders"
    styles = ["thin", "medium", "thick", "dashed", "dotted", "double", "hair", "mediumDashed", "dashDot", "mediumDashDot", "dashDotDot", "slantDashDot"]
    for i, s in enumerate(styles):
        row = 2 + i * 2
        ws.cell(row=row, column=1, value=s)
        c = ws.cell(row=row, column=3)
        side = Side(style=s, color="000000")
        c.border = Border(left=side, right=side, top=side, bottom=side)
    colored = Side(style="medium", color="C00000")
    ws["E2"].border = Border(left=colored, right=colored, top=colored, bottom=colored)
    ws["E4"].border = Border(diagonal=Side(style="thin", color="0070C0"), diagonalUp=True, diagonalDown=True)
    # Shared edge: thick wins over thin.
    ws["G2"].border = Border(right=Side(style="thin"))
    ws["H2"].border = Border(left=Side(style="thick"))
    for r in range(6, 12):
        for c in range(5, 9):
            ws.cell(row=r, column=c).border = Border(
                left=Side(style="thin", color="808080"),
                right=Side(style="thin", color="808080"),
                top=Side(style="thin", color="808080"),
                bottom=Side(style="thin", color="808080"),
            )
            ws.cell(row=r, column=c, value=(r - 5) * (c - 4))
    ws.sheet_view.showGridLines = False
    save(wb, "borders.xlsx")


def merges() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Merges"
    ws.merge_cells("A1:D1")
    ws["A1"] = "Merged title centred across A1:D1"
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor="DDEBF7")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A3:A8")
    ws["A3"] = "Vertical merge"
    ws["A3"].alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws["A3"].fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin")
    for r in range(3, 9):
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in range(3, 9):
        ws.cell(row=r, column=2, value=f"Row {r}")
    ws.merge_cells("C3:D4")
    ws["C3"] = "2x2"
    ws["C3"].alignment = Alignment(horizontal="right", vertical="bottom")
    ws.merge_cells("F2:H40")
    ws["F2"] = "Tall merge: text stays visible while its top row scrolls out of view"
    ws["F2"].alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws["F2"].fill = PatternFill("solid", fgColor="E2EFDA")
    save(wb, "merges.xlsx")


def text() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Text"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws["A1"] = "Wrapped text that is long enough to need several lines in this cell"
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 60
    ws["B1"] = "줄바꿈이 필요한 한글 문장입니다. 셀 너비에 맞춰 글자 단위로 끊어집니다."
    ws["B1"].alignment = Alignment(wrap_text=True)
    ws["C1"] = "Line one\nLine two\nLine three"
    ws["C1"].alignment = Alignment(wrap_text=True)
    ws["A3"] = "Left overflow runs across empty cells"
    ws["E4"] = "Right aligned overflow goes left"
    ws["E4"].alignment = Alignment(horizontal="right")
    ws["C5"] = "Centred overflow spreads both ways"
    ws["C5"].alignment = Alignment(horizontal="center")
    ws["A6"] = "Blocked by the next cell"
    ws["B6"] = "x"
    ws["A8"] = "Indent 2"
    ws["A8"].alignment = Alignment(indent=2)
    ws["A9"] = "Rotated 45"
    ws["A9"].alignment = Alignment(text_rotation=45)
    ws["B9"] = "Rotated 90"
    ws["B9"].alignment = Alignment(text_rotation=90)
    ws["C9"] = "Vertical"
    ws["C9"].alignment = Alignment(text_rotation=255)
    ws.row_dimensions[9].height = 80
    ws["A11"] = "Shrink to fit this long text"
    ws["A11"].alignment = Alignment(shrink_to_fit=True)
    ws["A12"] = "*"
    ws["A12"].alignment = Alignment(horizontal="fill")
    ws["A13"] = "Top"
    ws["A13"].alignment = Alignment(vertical="top")
    ws["B13"] = "Middle"
    ws["B13"].alignment = Alignment(vertical="center")
    ws["C13"] = "Bottom"
    ws.row_dimensions[13].height = 40
    ws["A14"] = "x"
    ws["B14"] = "Superscript"
    ws["B14"].font = Font(vertAlign="superscript")
    save(wb, "text.xlsx")


def freeze() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Freeze"
    ws.append(["ID"] + [f"Col {i}" for i in range(1, 30)])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for r in range(2, 301):
        ws.append([r - 1] + [random.randint(0, 9999) for _ in range(29)])
    ws.freeze_panes = "B2"
    ws.sheet_view.zoomScale = 125
    save(wb, "freeze.xlsx")


def freeze_rows_only() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RowsOnly"
    for r in range(1, 80):
        ws.append([f"r{r}c{c}" for c in range(1, 12)])
    ws.freeze_panes = "A3"
    save(wb, "freeze-rows-only.xlsx")


def png_bytes(w: int, h: int) -> bytes:
    from PIL import Image as PILImage, ImageDraw

    img = PILImage.new("RGB", (w, h), (68, 114, 196))
    d = ImageDraw.Draw(img)
    for i in range(0, w, 20):
        d.line([(i, 0), (i, h)], fill=(237, 125, 49), width=4)
    d.ellipse([w // 4, h // 4, w * 3 // 4, h * 3 // 4], fill=(255, 192, 0))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def notes_images() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws["A1"] = "Hover me"
    ws["A1"].comment = Comment("A plain note.\nSecond line.", "Reviewer")
    ws["C3"] = "Wide note"
    c = Comment("A wider note box set through the VML size.", "Reviewer")
    c.width = 260
    c.height = 60
    ws["C3"].comment = c
    ws["A6"] = "Image below"
    img = Image(io.BytesIO(png_bytes(160, 100)))
    ws.add_image(img, "B7")
    save(wb, "notes-images.xlsx")


def charts() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Month", "North", "South", "East"])
    for m in ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]:
        ws.append([m, random.randint(20, 90), random.randint(20, 90), random.randint(20, 90)])
    data = Reference(ws, min_col=2, min_row=1, max_col=4, max_row=7)
    cats = Reference(ws, min_col=1, min_row=2, max_row=7)

    def place(chart, anchor, title):
        chart.title = title
        chart.width = 12
        chart.height = 7
        ws.add_chart(chart, anchor)

    bar = BarChart()
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    place(bar, "F1", "Clustered column")

    stacked = BarChart()
    stacked.grouping = "stacked"
    stacked.overlap = 100
    stacked.add_data(data, titles_from_data=True)
    stacked.set_categories(cats)
    place(stacked, "N1", "Stacked column")

    horiz = BarChart()
    horiz.type = "bar"
    horiz.grouping = "percentStacked"
    horiz.overlap = 100
    horiz.add_data(data, titles_from_data=True)
    horiz.set_categories(cats)
    place(horiz, "F16", "100% stacked bar")

    line = LineChart()
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    place(line, "N16", "Line")

    pie = PieChart()
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    pie.set_categories(cats)
    place(pie, "F31", "Pie")

    dough = DoughnutChart()
    dough.add_data(Reference(ws, min_col=3, min_row=1, max_row=7), titles_from_data=True)
    dough.set_categories(cats)
    place(dough, "N31", "Doughnut")

    area = AreaChart()
    area.grouping = "stacked"
    area.add_data(data, titles_from_data=True)
    area.set_categories(cats)
    place(area, "F46", "Stacked area")

    scatter = ScatterChart()
    scatter.style = 13
    xs = Reference(ws, min_col=2, min_row=2, max_row=7)
    ys = Reference(ws, min_col=3, min_row=1, max_row=7)
    s = Series(ys, xs, title_from_data=True)
    s.marker.symbol = "circle"
    s.graphicalProperties.line.noFill = True
    scatter.series.append(s)
    place(scatter, "N46", "Scatter")

    combo = BarChart()
    combo.add_data(Reference(ws, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    combo.set_categories(cats)
    combo_line = LineChart()
    combo_line.add_data(Reference(ws, min_col=4, min_row=1, max_row=7), titles_from_data=True)
    combo_line.y_axis.axId = 200
    combo_line.y_axis.crosses = "max"
    combo += combo_line
    place(combo, "F61", "Combo with secondary axis")
    save(wb, "charts.xlsx")


def conditional() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Conditional"
    ws.append(["Scale", "Bar", "Icons", "CellIs", "Top/Avg", "Text"])
    words = ["apple", "banana", "cherry", "apple", "date", "banana", "fig", "grape", "apple", "kiwi"]
    for i in range(10):
        ws.append([i * 10, i * 10, i * 10, i * 10 - 30, random.randint(0, 100), words[i]])
    ws.conditional_formatting.add("A2:A11", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    ws.conditional_formatting.add("B2:B11", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    ws.conditional_formatting.add("C2:C11", IconSetRule("3Arrows", "percent", [0, 33, 67]))
    red = DifferentialStyle(font=Font(color="9C0006"), fill=PatternFill(bgColor="FFC7CE"))
    ws.conditional_formatting.add("D2:D11", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006"), fill=PatternFill(bgColor="FFC7CE")))
    ws.conditional_formatting.add("D2:D11", CellIsRule(operator="between", formula=["10", "40"], font=Font(bold=True)))
    top = Rule(type="top10", rank=3, dxf=DifferentialStyle(fill=PatternFill(bgColor="C6EFCE")))
    ws.conditional_formatting.add("E2:E11", top)
    above = Rule(type="aboveAverage", aboveAverage=True, dxf=DifferentialStyle(font=Font(bold=True)))
    ws.conditional_formatting.add("E2:E11", above)
    dup = Rule(type="duplicateValues", dxf=red)
    ws.conditional_formatting.add("F2:F11", dup)
    ws.conditional_formatting.add("A13:B13", FormulaRule(formula=["A13>0"], fill=PatternFill(bgColor="FFFF00")))
    ws["A13"] = 5
    ws["B13"] = "expression rules are not evaluated"
    save(wb, "conditional.xlsx")


def large() -> None:
    from openpyxl import Workbook as WB

    wb = WB(write_only=True)
    ws = wb.create_sheet("Large")
    ws.append([f"Col {i}" for i in range(1, 21)])
    for r in range(100_000):
        ws.append([r, f"row {r}", r * 1.5, r % 7 == 0] + [random.randint(0, 100000) for _ in range(16)])
    (OUT / "perf").mkdir(parents=True, exist_ok=True)
    path = OUT / "perf" / "large-100k.xlsx"
    wb.save(path)
    print(path.relative_to(OUT.parent.parent))


def resize() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resize"
    ws.append(["Name", "Endpoint", "Change"])
    notes = [
        "Summary: status and situation changed.\n[Request] key renamed.\n[Response] same.",
        "One short line",
        "A long sentence without line breaks that needs several lines once the column is narrow enough to wrap it.",
    ]
    for i, note in enumerate(notes, start=2):
        ws.append([f"item {i - 1}", f"/api/v1/admin/items/{i}/a-very-long-endpoint-name", note])
        ws.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 18
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 40
    save(wb, "resize.xlsx")


if __name__ == "__main__":
    basic()
    number_formats()
    borders()
    merges()
    text()
    freeze()
    freeze_rows_only()
    resize()
    notes_images()
    charts()
    conditional()
    if "--large" in sys.argv:
        large()
